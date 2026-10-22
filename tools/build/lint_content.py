"""Lint v4 course content for the rules in AGENTS.md.

Usage:
    python tools/build/lint_content.py              # all v4 content
    python tools/build/lint_content.py Day12 Day13  # only these folders (and their specs)

Checks text in .md, .txt, .ipynb, specs and the text inside .pptx, .docx and .pdf files for:
- em dashes
- other AI vendors' assistants or instruction files
- the retired v1 default model name
- hard-coded API keys
The archive and FDE-Mindset-Source folders are not checked.
"""
import glob
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RULES = [
    ("em dash", re.compile("—")),
    ("other AI vendor", re.compile(r"\b(Claude|Anthropic|Copilot|Cursor|Gemini|Windsurf|CLAUDE\.md)\b")),
    ("retired default model", re.compile(r"gpt-4o-mini")),
    ("possible API key", re.compile(r"sk-[A-Za-z0-9_-]{20,}")),
]


def text_of(path):
    if path.endswith((".md", ".txt", ".py", ".ipynb", ".yml", ".yaml", ".json", ".js", ".jsx", ".toml", ".cfg")):
        return open(path, encoding="utf-8", errors="ignore").read()
    if path.endswith(".pptx"):
        from pptx import Presentation
        out = []
        for s in Presentation(path).slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    out.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    out += [c.text for r in sh.table.rows for c in r.cells]
        return "\n".join(out)
    if path.endswith(".docx"):
        import docx
        d = docx.Document(path)
        return "\n".join(p.text for p in d.paragraphs) + "\n".join(
            c.text for t in d.tables for r in t.rows for c in r.cells)
    if path.endswith(".pdf"):
        import pdfplumber
        with pdfplumber.open(path) as p:
            return "\n".join(pg.extract_text() or "" for pg in p.pages)
    return ""


def main():
    args = sys.argv[1:]
    files = []
    for pattern in ("Day*/**/*", "shared/**/*", "docs/**/*", "*.md", "tools/build/specs/*.py"):
        files += [f for f in glob.glob(os.path.join(ROOT, pattern), recursive=True) if os.path.isfile(f)]
    files = [f for f in files if "node_modules" not in f and ".ipynb_checkpoints" not in f]
    if args:
        files = [f for f in files if any(a in os.path.relpath(f, ROOT) or
                                         os.path.basename(f).startswith(f"day{a[3:5]}") for a in args)]
    problems = 0
    for f in sorted(set(files)):
        rel = os.path.relpath(f, ROOT)
        if rel == "tools/build/lint_content.py":
            continue
        for i, line in enumerate(text_of(f).splitlines(), 1):
            for name, rx in RULES:
                if rx.search(line):
                    problems += 1
                    print(f"{rel}:{i}: {name}: {line.strip()[:110]}")
    print(f"{problems} problem(s) found.")
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
