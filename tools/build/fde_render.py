"""Renderer for FDE Program v4 day materials.

Turns a day spec (a Python module in tools/build/specs/) into the deck (.pptx),
pre-read, post-read and hands-on PDFs, workbook (.docx) and Colab notebook,
using the same visual style as the v1 materials.

Inline markup in any text field: **bold**. Use "NEW (v4)" in text where a topic is new.
"""
import json
import re

# ---------------------------------------------------------------- palette
DARK = "0E3B43"      # headings, body text
DARK_BG = "0B5D63"   # dark slide background
TEAL = "0FA3A3"
ORANGE = "FF8A3D"
ORANGE_DOC = "FF7A2F"
CARD = "E3F6F5"
CARD2 = "EAF7F5"
BODY = "1F4E55"
MUTED = "3D7C82"
CODE_BG = "0F2C33"
CODE_FG = "EAF7F5"
LIGHT_ON_DARK = "CDEFEC"
GRID = "E5E7EB"

AUTHOR = "Ahmed"
LINKEDIN = "linkedin.com/in/ahmed-al-imaad"
LINKEDIN_URL = "https://www.linkedin.com/in/ahmed-al-imaad"

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def split_bold(text):
    """Return [(chunk, is_bold), ...] for **bold** markup."""
    out, pos = [], 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], False))
        out.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], False))
    return out or [("", False)]


# ================================================================ DECK
def build_deck(spec, path):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    def rgb(h):
        return RGBColor.from_string(h)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(5.625)
    blank = prs.slide_layouts[6]
    day_label = f"Day {spec.DAY}  |  {spec.TITLE}"
    counter = {"n": 0}

    def bg(slide, color):
        f = slide.background.fill
        f.solid()
        f.fore_color.rgb = rgb(color)

    def text(slide, x, y, w, h, value, size=14, color=DARK, bold=False, font="Calibri",
             align=None, anchor=MSO_ANCHOR.TOP, italic=False, spacing=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        lines = value if isinstance(value, list) else [value]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if align:
                p.alignment = align
            for chunk, b in split_bold(line):
                r = p.add_run()
                r.text = chunk
                r.font.size = Pt(size)
                r.font.name = font
                r.font.bold = bold or b
                r.font.italic = italic
                r.font.color.rgb = rgb(color)
            if spacing:
                p.space_after = Pt(spacing)
        return tb

    def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
        s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(color)
        s.line.fill.background()
        s.shadow.inherit = False
        return s

    def badge(slide, x, y, n, color, d=0.42):
        rect(slide, x, y, d, d, color, MSO_SHAPE.OVAL)
        text(slide, x, y, d, d, str(n), 12, "FFFFFF", True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def new_tag(slide):
        rect(slide, 8.35, 0.12, 1.15, 0.3, ORANGE, MSO_SHAPE.ROUNDED_RECTANGLE)
        text(slide, 8.35, 0.12, 1.15, 0.3, "NEW (v4)", 11, "FFFFFF", True,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def content_slide(s):
        slide = prs.slides.add_slide(blank)
        counter["n"] += 1
        bg(slide, "FFFFFF")
        title = s.get("title", "")
        text(slide, 0.5, 0.28, 7.8 if s.get("new") else 9.0, 0.7, title, 26, DARK, True, "Cambria",
             anchor=MSO_ANCHOR.MIDDLE)
        if s.get("new"):
            new_tag(slide)
        text(slide, 0.5, 5.2, 4.0, 0.3, day_label, 9, TEAL, anchor=MSO_ANCHOR.MIDDLE)
        if s.get("greet"):
            text(slide, 4.6, 5.2, 4.2, 0.3, s["greet"], 9, ORANGE, italic=True,
                 align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        text(slide, 9.0, 5.2, 0.5, 0.3, str(counter["n"]), 9, MUTED, align=PP_ALIGN.RIGHT,
             anchor=MSO_ANCHOR.MIDDLE)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
        return slide

    def note_line(slide, s, y):
        if s.get("note"):
            text(slide, 0.5, y, 9.0, 0.5, s["note"], 13, TEAL, italic=True)

    def fit(t, base, length_limits):
        n = len(t)
        for limit, size in length_limits:
            if n > limit:
                return size
        return base

    def cards(slide, items, top=1.15, bottom=4.85):
        n = len(items)
        cols = 2 if n in (2, 4) else 3
        rows = (n + cols - 1) // cols
        gap = 0.25
        w = (9.0 - gap * (cols - 1)) / cols
        h = (bottom - top - gap * (rows - 1)) / rows
        for i, it in enumerate(items):
            title, body = it[0], it[1] if len(it) > 1 else ""
            r, c = divmod(i, cols)
            x, y = 0.5 + c * (w + gap), top + r * (h + gap)
            rect(slide, x, y, w, h, CARD)
            badge(slide, x + 0.2, y + 0.2, i + 1, ORANGE if i % 2 == 0 else TEAL)
            text(slide, x + 0.75, y + 0.2, w - 0.95, 0.42, title, fit(title, 15, [(34, 12), (24, 13)]),
                 DARK, True, anchor=MSO_ANCHOR.MIDDLE)
            text(slide, x + 0.2, y + 0.72, w - 0.4, h - 0.85, body,
                 fit(body, 14, [(170, 10.5), (120, 11.5), (80, 12.5)]), BODY)

    def bullets(slide, items, x, y, w, h, size=16):
        size = fit(" ".join(items), size, [(650, 11.5), (480, 12.5), (360, 14)])
        text(slide, x, y, w, h, [f"•  {b}" for b in items], size, DARK, spacing=6)

    def code_box(slide, code, x, y, w, h):
        rect(slide, x, y, w, h, CODE_BG)
        lines = code.strip("\n").split("\n")
        longest = max(len(l) for l in lines)
        by_width = 0.88 * (w - 0.4) * 72 / (0.6 * max(longest, 1))  # Courier glyph is 0.6 em; keep a margin
        by_height = (h - 0.3) * 72 / (1.2 * max(len(lines), 1))
        size = max(7, min(14, by_width, by_height))
        tb = text(slide, x + 0.2, y + 0.15, w - 0.4, h - 0.3, lines, size, CODE_FG, font="Courier New")
        tb.text_frame.word_wrap = False                               # code lines never wrap

    def table(slide, header, rows, top=1.2):
        nrows, ncols = len(rows) + 1, len(header)
        rh = min(0.42, 3.6 / nrows)
        gs = slide.shapes.add_table(nrows, ncols, Inches(0.5), Inches(top), Inches(9.0), Inches(rh * nrows))
        t = gs.table
        size = 11 if nrows <= 7 else 10 if nrows <= 9 else 9
        for ci in range(ncols):
            t.columns[ci].width = Inches(9.0 / ncols)
        for ri in range(nrows):
            vals = header if ri == 0 else rows[ri - 1]
            for ci in range(ncols):
                cell = t.cell(ri, ci)
                cell.fill.solid()
                cell.fill.fore_color.rgb = rgb(TEAL if ri == 0 else ("FFFFFF" if ri % 2 else CARD2))
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                for chunk, b in split_bold(str(vals[ci])):
                    r = p.add_run()
                    r.text = chunk
                    r.font.size = Pt(size)
                    r.font.name = "Calibri"
                    r.font.bold = (ri == 0) or b
                    r.font.color.rgb = rgb("FFFFFF" if ri == 0 else DARK)
        return top + rh * nrows

    def dark_slide(title, sub=None, kicker=None, big=None):
        slide = prs.slides.add_slide(blank)
        counter["n"] += 1
        bg(slide, DARK_BG)
        y = 1.0
        if big:
            text(slide, 0.6, 0.95, 4, 0.9, big, 40, ORANGE, True, "Cambria")
            y = 1.95
        if kicker:
            text(slide, 0.7, 1.0, 3, 0.4, kicker, 14, ORANGE, True, anchor=MSO_ANCHOR.MIDDLE)
            y = 1.4
        text(slide, 0.6 if big else 0.7, y, 8.6, 1.2, title, 34 if not big else 30, "FFFFFF", True, "Cambria")
        if sub:
            text(slide, 0.6 if big else 0.7, y + (1.25 if not big else 0.75), 8.6, 0.8, sub, 17 if not big else 14,
                 LIGHT_ON_DARK)
        return slide

    # ---- title slide
    slide = dark_slide(spec.TITLE, getattr(spec, "SUBTITLE", None), kicker=f"DAY {spec.DAY}")
    for i, col in enumerate([ORANGE, TEAL, "FFFFFF"]):
        rect(slide, 0.7 + i * 0.28, 3.6, 0.18, 0.18, col, MSO_SHAPE.OVAL)
    meta = f"{spec.DATE}  |  {spec.TRACK} track"
    text(slide, 0.7, 3.95, 8.0, 0.3, meta, 12, LIGHT_ON_DARK)
    text(slide, 0.7, 4.45, 4, 0.25, f"Prepared by {AUTHOR}", 11, "FFFFFF", True)
    text(slide, 0.7, 4.72, 4, 0.25, LINKEDIN, 11, LIGHT_ON_DARK)

    section_no = 0
    for s in spec.DECK:
        t = s["type"]
        if t == "section":
            section_no += 1
            sl = dark_slide(s["title"], s.get("sub"), big=f"{section_no:02d}")
            if s.get("new"):
                rect(sl, 8.35, 0.25, 1.15, 0.3, ORANGE, MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 8.35, 0.25, 1.15, 0.3, "NEW (v4)", 11, "FFFFFF", True, align=PP_ALIGN.CENTER,
                     anchor=MSO_ANCHOR.MIDDLE)
            continue
        slide = content_slide(s)
        if t == "cards":
            top = 1.15
            if s.get("lead"):
                text(slide, 0.5, 1.05, 9.0, 0.35, s["lead"], 13, TEAL)
                top = 1.5
            cards(slide, s["items"], top=top, bottom=4.45 if s.get("note") else 4.85)
            note_line(slide, s, 4.6)
        elif t == "flow":
            steps = s["items"]
            n = len(steps)
            gap = 0.35
            w = (9.0 - gap * (n - 1)) / n
            for i, (ti, body) in enumerate(steps):
                x = 0.5 + i * (w + gap)
                rect(slide, x, 1.5, w, 2.5, CARD)
                badge(slide, x + w / 2 - 0.25, 1.7, i + 1, ORANGE if i % 2 == 0 else TEAL, 0.5)
                text(slide, x + 0.1, 2.35, w - 0.2, 0.45, ti, fit(ti, 14, [(18, 12)]), DARK, True,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(slide, x + 0.12, 2.8, w - 0.24, 1.1, body, fit(body, 13, [(70, 11), (50, 12)]), BODY,
                     align=PP_ALIGN.CENTER)
                if i < n - 1:
                    rect(slide, x + w + 0.09, 2.6, 0.17, 0.3, TEAL, MSO_SHAPE.CHEVRON)
            note_line(slide, s, 4.35)
        elif t == "table":
            end = table(slide, s["header"], s["rows"])
            note_line(slide, s, min(end + 0.15, 4.65))
        elif t == "bullets_code":
            bullets(slide, s["bullets"], 0.5, 1.25, 4.1, 3.6)
            code_box(slide, s["code"], 4.85, 1.2, 4.65, 3.6 if not s.get("note") else 3.2)
            note_line(slide, s, 4.55)
        elif t == "code":
            code_box(slide, s["code"], 0.5, 1.15, 9.0, 3.4 if s.get("note") else 3.8)
            note_line(slide, s, 4.65)
        elif t == "compare":
            for i, side in enumerate([s["left"], s["right"]]):
                x = 0.5 + i * 4.6
                rect(slide, x, 1.2, 4.4, 3.1 if s.get("note") else 3.5, CARD)
                text(slide, x + 0.25, 1.35, 3.9, 0.4, side["title"], 15, ORANGE if i == 0 else TEAL, True)
                size = fit(" ".join(side["bullets"]), 14, [(380, 11), (260, 12), (180, 13)])
                text(slide, x + 0.25, 1.85, 3.9, 2.4, [f"•  {b}" for b in side["bullets"]], size, DARK,
                     spacing=4)
            note_line(slide, s, 4.5)
        elif t == "bullets":
            bullets(slide, s["bullets"], 0.5, 1.25, 9.0, 3.3 if s.get("note") else 3.7)
            note_line(slide, s, 4.6)
        elif t == "callout":
            rect(slide, 0.5, 1.3, 9.0, 1.9, TEAL)
            text(slide, 0.8, 1.45, 8.4, 1.6, s["text"], fit(s["text"], 20, [(220, 15), (160, 17)]), "FFFFFF",
                 True, "Cambria", anchor=MSO_ANCHOR.MIDDLE)
            if s.get("sub"):
                text(slide, 0.5, 3.45, 9.0, 1.2, s["sub"], 13, DARK)
            note_line(slide, s, 4.6)
        else:
            raise ValueError(f"Unknown slide type {t}")

    # ---- closing slides
    slide = dark_slide("Connect with me", "Questions after class, feedback, or want to share what you built today? "
                       "Reach out on LinkedIn.")
    rect(slide, 0.7, 3.4, 4.5, 0.55, ORANGE, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(slide, 0.9, 3.4, 4.2, 0.55, f"**{AUTHOR}**   {LINKEDIN}", 13, "FFFFFF", anchor=MSO_ANCHOR.MIDDLE)
    slide = dark_slide("Thank you", getattr(spec, "NEXT", "See you tomorrow."))
    text(slide, 4.6, 5.2, 4.9, 0.3, getattr(spec, "CLOSING_GREET",
         "Otsukaresama deshita (お疲れ様でした): Thank you for your hard work today"), 9, LIGHT_ON_DARK,
         italic=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    prs.save(path)


# ================================================================ PDFs
def _pdf_styles():
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    c = HexColor
    return {
        "kicker": ParagraphStyle("k", fontName="Helvetica-Bold", fontSize=10, textColor=c("#" + ORANGE_DOC),
                                 spaceAfter=6),
        "title": ParagraphStyle("t", fontName="Times-Bold", fontSize=28, leading=33, textColor=c("#" + DARK),
                                spaceAfter=8),
        "lede": ParagraphStyle("l", fontName="Helvetica", fontSize=12.5, leading=16, textColor=c("#" + MUTED),
                               spaceAfter=10),
        "by": ParagraphStyle("b", fontName="Helvetica", fontSize=9, textColor=c("#" + TEAL), spaceAfter=14),
        "h2": ParagraphStyle("h2", fontName="Times-Bold", fontSize=17, leading=21, textColor=c("#" + DARK),
                             spaceBefore=12, spaceAfter=6),
        "h3": ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=c("#" + TEAL),
                             spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("p", fontName="Helvetica", fontSize=10.5, leading=15, textColor=c("#" + DARK),
                               spaceAfter=6),
        "meta": ParagraphStyle("m", fontName="Helvetica", fontSize=9, leading=12, textColor=c("#" + MUTED),
                               spaceAfter=4),
        "bullet": ParagraphStyle("bl", fontName="Helvetica", fontSize=10.5, leading=15,
                                 textColor=c("#" + DARK), leftIndent=14, bulletIndent=4, spaceAfter=2),
        "code": ParagraphStyle("c", fontName="Courier", fontSize=9.5, leading=12.5, textColor=c("#" + CODE_FG)),
        "box": ParagraphStyle("x", fontName="Helvetica", fontSize=10, leading=14, textColor=c("#" + DARK)),
        "newtag": ParagraphStyle("n", fontName="Helvetica-Bold", fontSize=9, textColor=c("#" + ORANGE_DOC)),
    }


CJK_RE = re.compile(r"[　-ヿ㐀-鿿＀-￯]+")
_CJK = {"font": None}


def _cjk_font():
    """Register IPAGothic (as in v1) for Japanese text, if installed."""
    if _CJK["font"] is None:
        _CJK["font"] = ""
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        for p in ("/usr/share/fonts/opentype/ipafont-gothic/ipag.ttf",
                  "/usr/share/fonts/truetype/fonts-japanese-gothic.ttf"):
            try:
                pdfmetrics.registerFont(TTFont("IPAGothic", p))
                _CJK["font"] = "IPAGothic"
                break
            except Exception:
                continue
    return _CJK["font"]


def _md(text):
    """Escape for ReportLab and convert **bold**, `code` and Japanese text."""
    t = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    t = BOLD_RE.sub(r"<b>\1</b>", t)
    t = re.sub(r"`([^`]+)`", r'<font face="Courier">\1</font>', t)
    if CJK_RE.search(t) and _cjk_font():
        t = CJK_RE.sub(lambda m: f'<font name="IPAGothic">{m.group(0)}</font>', t)
    return t


def _box(flow, text, st, fill=CARD):
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import Paragraph, Table, TableStyle
    tb = Table([[Paragraph(_md(text), st["box"])]], colWidths=[481.9])
    tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), HexColor("#" + fill)),
                            ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                            ("TOPPADDING", (0, 0), (-1, -1), 9), ("BOTTOMPADDING", (0, 0), (-1, -1), 9)]))
    flow.append(tb)


def _code(flow, code, st):
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import Preformatted, Table, TableStyle
    tb = Table([[Preformatted(code.strip("\n"), st["code"])]], colWidths=[481.9])
    tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), HexColor("#" + CODE_BG)),
                            ("LEFTPADDING", (0, 0), (-1, -1), 10), ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                            ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    flow.append(tb)


def _table(flow, header, rows, st):
    from reportlab.lib.colors import HexColor
    from reportlab.platypus import Paragraph, Table, TableStyle
    data = [[Paragraph(f"<b>{_md(str(h))}</b>", st["box"]) for h in header]]
    data += [[Paragraph(_md(str(v)), st["box"]) for v in r] for r in rows]
    tb = Table(data, colWidths=[481.9 / len(header)] * len(header), repeatRows=1)
    tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), HexColor("#" + CARD)),
                            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#FFFFFF"), HexColor("#" + CARD2)]),
                            ("GRID", (0, 0), (-1, -1), 0.4, HexColor("#9ED9D4")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    flow.append(tb)


def _render_blocks(flow, blocks, st):
    """Generic section renderer: list of dicts with h, paras, bullets, code, table, box, sub."""
    from reportlab.platypus import Paragraph, Spacer
    for sec in blocks:
        if sec.get("h"):
            h = _md(sec["h"])
            if sec.get("new"):
                h += '  <font name="Helvetica-Bold" size="9" color="#FF7A2F">NEW (v4)</font>'
            flow.append(Paragraph(h, st["h2"]))
        if sec.get("sub"):
            flow.append(Paragraph(_md(sec["sub"]), st["h3"]))
        for p in sec.get("paras", []):
            flow.append(Paragraph(_md(p), st["body"]))
        for b in sec.get("bullets", []):
            flow.append(Paragraph(_md(b), st["bullet"], bulletText="•"))
        if sec.get("table"):
            _table(flow, sec["table"][0], sec["table"][1:], st)
            flow.append(Spacer(1, 6))
        if sec.get("code"):
            _code(flow, sec["code"], st)
            flow.append(Spacer(1, 6))
        if sec.get("box"):
            flow.append(Spacer(1, 4))
            _box(flow, sec["box"], st)
            flow.append(Spacer(1, 6))
        for p in sec.get("after", []):
            flow.append(Paragraph(_md(p), st["body"]))


def _doc(path, title, footer, flow):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(HexColor("#" + MUTED))
        canvas.drawString(56.7, 28, footer)
        canvas.drawRightString(A4[0] - 56.7, 28, f"Page {doc.page}")
        canvas.restoreState()

    d = SimpleDocTemplate(path, pagesize=A4, leftMargin=56.7, rightMargin=56.7, topMargin=50, bottomMargin=56,
                          title=title, author=AUTHOR)
    d.build(flow, onFirstPage=on_page, onLaterPages=on_page)


def _header(flow, st, kicker, title, lede):
    from reportlab.platypus import Paragraph
    flow.append(Paragraph(kicker, st["kicker"]))
    flow.append(Paragraph(_md(title), st["title"]))
    if lede:
        flow.append(Paragraph(_md(lede), st["lede"]))
    flow.append(Paragraph(f'Prepared by <b>{AUTHOR}</b> | <a href="{LINKEDIN_URL}" color="#0FA3A3">{LINKEDIN}</a>',
                          st["by"]))


def build_read(spec, kind, path):
    """kind: 'pre' or 'post'."""
    st = _pdf_styles()
    data = spec.PREREAD if kind == "pre" else spec.POSTREAD
    label = "PRE-READ" if kind == "pre" else "POST-READ"
    flow = []
    _header(flow, st, f"DAY {spec.DAY} | {label}", data["title"], data.get("lede"))
    _render_blocks(flow, data["sections"], st)
    _doc(path, f"Day {spec.DAY}  |  {label.title()}", f"Day {spec.DAY}  |  {label.title()}", flow)


def build_handson(spec, path):
    from reportlab.platypus import KeepTogether, Paragraph, Spacer
    st = _pdf_styles()
    h = spec.HANDSON
    flow = []
    _header(flow, st, f"DAY {spec.DAY} | HANDS-ON ACTIVITIES", f"{spec.TITLE}: Problem Statements", h.get("lede"))
    _box(flow, h.get("how", "**How to work:** read the problem, write the steps in plain English first, then write "
                     "the code. Test with the sample run, then try your own inputs, including wrong ones."), st)
    flow.append(Spacer(1, 8))
    for i, p in enumerate(h["problems"], 1):
        head = f"Problem {i}: {_md(p['title'])}"
        if p.get("new"):
            head += '  <font name="Helvetica-Bold" size="9" color="#FF7A2F">NEW (v4)</font>'
        block = [Paragraph(head, st["h2"]),
                 Paragraph(f"<b>Level:</b> {p['level']} | <b>Concepts:</b> {_md(p['concepts'])}", st["meta"])]
        flow.append(KeepTogether(block))
        for para in (p["text"] if isinstance(p["text"], list) else [p["text"]]):
            flow.append(Paragraph(_md(para), st["body"]))
        for b in p.get("bullets", []):
            flow.append(Paragraph(_md(b), st["bullet"], bulletText="•"))
        if p.get("sample"):
            flow.append(Paragraph("Sample run / expected result", st["h3"]))
            _code(flow, p["sample"], st)
        if p.get("hints"):
            flow.append(Paragraph("Hints", st["h3"]))
            for hint in p["hints"]:
                flow.append(Paragraph(_md(hint), st["bullet"], bulletText="•"))
    _doc(path, f"Day {spec.DAY}  |  Hands-on activities", f"Day {spec.DAY}  |  Hands-on activities", flow)


# ================================================================ WORKBOOK
def build_workbook(spec, path):
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    def para(t, size=11, color=DARK, bold=False, font="Calibri", before=0, after=80, keep=False, mono=False):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(before / 20), Pt(after / 20)
        pf.keep_with_next = keep
        for chunk, b in split_bold(t):
            r = p.add_run(chunk)
            r.font.size = Pt(size)
            r.font.name = "Courier New" if mono else font
            r._element.rPr.rFonts.set(qn("w:eastAsia"), r.font.name)
            r.font.bold = bold or b
            r.font.color.rgb = RGBColor.from_string(color)
        return p

    def border_bottom(p):
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        b = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "4"), ("w:color", TEAL)):
            b.set(qn(k), v)
        bdr.append(b)
        pPr.append(bdr)

    def answer_box(lines=1, prefill=None):
        t = doc.add_table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{side}")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:color", "9ED9D4")):
                e.set(qn(k), v)
            borders.append(e)
        tcPr.append(borders)
        shd = OxmlElement("w:shd")
        for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", "FAFDFD")):
            shd.set(qn(k), v)
        tcPr.append(shd)
        tr = t.rows[0]._tr
        trPr = tr.get_or_add_trPr()
        h = OxmlElement("w:trHeight")
        h.set(qn("w:val"), str(300 * lines))
        h.set(qn("w:hRule"), "atLeast")
        trPr.append(h)
        p = cell.paragraphs[0]
        if prefill:
            for i, line in enumerate(prefill.strip("\n").split("\n")):
                r = p.add_run(line)
                r.font.name = "Courier New"
                r.font.size = Pt(10)
                if i < len(prefill.strip("\n").split("\n")) - 1:
                    r.add_break(WD_BREAK.LINE)
        para("", after=80)

    para(f"DAY {spec.DAY}  |  WORKBOOK", 10, ORANGE_DOC, True, after=60)
    para(f"{spec.TITLE} Workbook", 22, DARK, True, "Cambria", after=80)
    para("Type your answers directly into the boxes in Word or Google Docs, or print it and write by hand.",
         11, MUTED, after=80)
    para(f"Prepared by **{AUTHOR}**  |  {LINKEDIN}", 9, MUTED, after=240)
    para("Name", keep=True, before=120)
    answer_box()

    for s in spec.WORKBOOK["sections"]:
        title = s["title"] + ("  NEW (v4)" if s.get("new") else "")
        p = para(title, 15, TEAL, True, "Cambria", before=280, after=120, keep=True)
        border_bottom(p)
        kind = s["kind"]
        if s.get("intro"):
            para(s["intro"], 10.5, MUTED, after=80)
        if kind == "reflection":
            for q in s.get("items", ["The idea that clicked best for me today:",
                                     "One thing I am still unsure about:",
                                     "One thing I will try in the hotel lab tonight:"]):
                para(q, keep=True, before=120)
                answer_box(3)
            continue
        for i, item in enumerate(s["items"], 1):
            if kind == "predict":
                para(f"Snippet {i}", bold=True, keep=True, before=120)
                para(item["code"] if isinstance(item, dict) else item, 10, DARK, mono=True, keep=True)
                para("Your answer:", 10, MUTED, keep=True)
                answer_box(2)
            elif kind == "code":
                q = item["q"] if isinstance(item, dict) else item
                para(f"{i}. {q}", keep=True, before=120)
                answer_box(8, item.get("starter") if isinstance(item, dict) else None)
            else:
                para(f"{i}. {item}", keep=True, before=120)
                answer_box(1 if kind == "blanks" else 4)
    doc.save(path)


# ================================================================ NOTEBOOK
GENAI_HELPERS = '''%%writefile genai_helpers.py
"""Shared helpers for the GenAI days.
llm()   -> calls OpenAI when a key is available, otherwise a small OFFLINE stand-in so every cell still runs.
embed() -> OpenAI embeddings online, a hashing-based stand-in offline (same API, lower quality).
usage() -> token counts of the last llm() call (estimated offline).
"""
import json, os, re
import numpy as np

def _get_key():
    if os.getenv("OFFLINE") == "1":                 # force offline demo mode (used by tools/check_notebooks.py)
        return None
    try:
        from google.colab import userdata          # Colab Secrets
        return userdata.get("OPENAI_API_KEY")
    except Exception:
        return os.getenv("OPENAI_API_KEY")

OPENAI_API_KEY = _get_key()
ONLINE = bool(OPENAI_API_KEY)
# Default model: confirm against the approved model list (docs/MODEL_GUIDE.md). Override with LLM_MODEL.
MODEL = os.getenv("LLM_MODEL", "gpt-6.1-sol")
EMBED_MODEL = os.getenv("EMBED_MODEL", "text-embedding-3-small")
if ONLINE:
    from openai import OpenAI
    client = OpenAI(api_key=OPENAI_API_KEY)

WORD = re.compile(r"[a-z0-9]+")
LAST_USAGE = {"input_tokens": 0, "output_tokens": 0}

_ENC = []

def count_tokens(text):
    if not _ENC:
        try:
            import tiktoken
            _ENC.append(tiktoken.get_encoding("o200k_base"))
        except Exception:
            _ENC.append(None)                        # no tiktoken or no download: use the estimate
    if _ENC[0] is not None:
        return len(_ENC[0].encode(text))
    return max(1, round(len(text) / 4))              # rough rule: about 4 characters per token

def _words(t):
    # crude normalisation so "expires" matches "expire"
    return {w[:-1] if w.endswith("s") and len(w) > 4 else w for w in WORD.findall(t.lower()) if len(w) > 3}

def offline_llm(prompt, system=""):
    """A deliberately simple stand-in. It is NOT intelligent; it lets the code paths run without a key."""
    low = prompt.lower()
    if "context:" in low and "question:" in low:
        ctx = re.split(r"(?i)context:", prompt, 1)[1]
        ctx, q = re.split(r"(?i)question:", ctx, 1)
        sents = [s.strip() for s in re.split(r"(?<=[.!?])\\s+", ctx.strip()) if s.strip()]
        best = max(sents, key=lambda s: len(_words(s) & _words(q)), default="")
        return best if best and _words(best) & _words(q) else "I don't know based on the provided context."
    if "classify" in low and "positive" in low:
        neg = sum(w in low for w in ["slow", "bad", "broken", "late", "angry", "worst", "refund"])
        pos = sum(w in low for w in ["great", "love", "fast", "thanks", "excellent", "happy"])
        return "negative" if neg > pos else "positive" if pos > neg else "neutral"
    if "classify" in low:                            # SmartDesk ticket categories
        last = low.strip().splitlines()[-1]
        for label, words in [("billing", ["bill", "charge", "refund", "invoice", "payment"]),
                             ("network", ["signal", "network", "slow", "outage", "data", "5g"]),
                             ("account", ["password", "login", "account", "sim", "plan"])]:
            if any(w in last for w in words):
                return label
        return "other"
    if "json" in low:
        text = prompt.strip().splitlines()[-1]
        return json.dumps({"summary": " ".join(text.split()[:12]), "offline": True})
    first = re.split(r"(?<=[.!?])\\s+", prompt.strip())[0]
    return f"[offline demo] {first[:160]}"

def llm(prompt, system="You are a helpful assistant.", temperature=0.2, model=None, max_tokens=None):
    if ONLINE:
        kwargs = {"max_tokens": max_tokens} if max_tokens else {}
        r = client.chat.completions.create(
            model=model or MODEL, temperature=temperature,
            messages=[{"role": "system", "content": system}, {"role": "user", "content": prompt}], **kwargs)
        LAST_USAGE.update(input_tokens=r.usage.prompt_tokens, output_tokens=r.usage.completion_tokens)
        return r.choices[0].message.content
    out = offline_llm(prompt, system)
    LAST_USAGE.update(input_tokens=count_tokens(system + prompt), output_tokens=count_tokens(out))
    return out

def usage():
    return dict(LAST_USAGE)

def embed(texts):
    """Return an (n, d) array of unit-length vectors."""
    if isinstance(texts, str):
        texts = [texts]
    if ONLINE:
        r = client.embeddings.create(model=EMBED_MODEL, input=texts)
        v = np.array([d.embedding for d in r.data])
    else:
        from sklearn.feature_extraction.text import HashingVectorizer
        hv = HashingVectorizer(n_features=1024, ngram_range=(1, 2), alternate_sign=False, norm=None)
        v = hv.transform([t.lower() for t in texts]).toarray()
    n = np.linalg.norm(v, axis=1, keepdims=True)
    return v / np.where(n == 0, 1, n)

print("Mode:", f"ONLINE ({MODEL})" if ONLINE else "OFFLINE demo (add OPENAI_API_KEY to use a real model)")
'''

GENAI_SETUP_MD = ("## Setup and API key\n"
                  "With an `OPENAI_API_KEY` the notebook calls a real model. Without one, it runs in "
                  "**offline demo mode** using a tiny stand-in, so every cell still works and you can study the code. "
                  "Offline scores show how the code works, not real model quality.\n\n"
                  "The model name comes from the `LLM_MODEL` environment variable. The default is the program's "
                  "approved default; confirm it against the approved model list in `docs/MODEL_GUIDE.md`.")


def build_notebook(spec, path):
    nbspec = spec.NOTEBOOK
    cells = []

    def md(s):
        cells.append({"cell_type": "markdown", "metadata": {}, "source": s.strip("\n")})

    def code(s, tags=None):
        c = {"cell_type": "code", "execution_count": None, "metadata": {}, "outputs": [], "source": s.strip("\n")}
        if tags:
            c["metadata"]["tags"] = tags
        cells.append(c)

    md(f"# Day {spec.DAY}: {spec.TITLE}\n**{nbspec.get('subtitle', getattr(spec, 'SUBTITLE', ''))}**\n\n"
       f"Prepared by **{AUTHOR}** | [{LINKEDIN}]({LINKEDIN_URL})\n\n---\n"
       "### How to use this notebook\n"
       "* Run a cell with **Shift + Enter**.\n"
       "* Some cells have two versions: **Option A (Google Colab)** is active, **Option B (VS Code / local)** is "
       "commented out.\n"
       "  * In Colab: run as is.\n"
       "  * In VS Code: comment Option A (select lines, press `Ctrl + /` or `Cmd + /`) and uncomment Option B.\n"
       "* Cells marked **TODO** are for you to complete.\n"
       "* Sections marked **NEW (v4)** were added in program v4.")
    for c in nbspec["cells"]:
        kind = c[0]
        if kind == "md":
            md(c[1])
        elif kind == "code":
            code(c[1], c[2] if len(c) > 2 else None)
        elif kind == "helpers":
            md(GENAI_SETUP_MD)
            code(GENAI_HELPERS)
            code("import importlib, genai_helpers\nimportlib.reload(genai_helpers)\n"
                 "from genai_helpers import llm, embed, usage, count_tokens, ONLINE, MODEL")
        else:
            raise ValueError(kind)
    md(f"---\n**{nbspec.get('sign_off', 'Ganbatte (頑張って): *Do your best*')}**\n\n"
       f"Prepared by **{AUTHOR}** | [{LINKEDIN}]({LINKEDIN_URL})")
    nb = {"cells": cells, "metadata": {"colab": {"provenance": []},
                                        "kernelspec": {"name": "python3", "display_name": "Python 3"},
                                        "language_info": {"name": "python"}},
          "nbformat": 4, "nbformat_minor": 0}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(nb, f, indent=1, ensure_ascii=False)
