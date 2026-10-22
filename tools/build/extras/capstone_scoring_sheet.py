"""Update Day36-40-Capstone/Capstone_Evaluator_Scoring_Sheet.xlsx for v4 (run once on the archive copy).
Rubric text uses v4 day numbers and baseline; a 'Baseline v4' sheet tracks the five NEW (v4) items per team."""
import copy
import os
import sys

import openpyxl
from openpyxl.styles import PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
sys.path.insert(0, HERE)
import capstone_content as C  # noqa: E402

PATH = os.path.join(ROOT, "Day36-40-Capstone", "Capstone_Evaluator_Scoring_Sheet.xlsx")
wb = openpyxl.load_workbook(PATH)

ins = wb["Instructions"]
ins["A1"] = "Capstone Evaluator Scoring Workbook (Days 36 to 40, Tue 15 Dec to Mon 21 Dec 2026)"
ins["A13"] = ("Grade bands: 85 to 100 Distinction, 70 to 84 Merit, 55 to 69 Pass, below 55 Needs resubmission "
              "(from the Capstone Guide).")
ins["A14"] = ("NEW (v4): also complete the 'Baseline v4' sheet (modernisation step, AGENTS.md, OWASP Agentic "
              "checklist, vulnerability report, token and cost report). Missing items lower the matching criterion.")
ins["A14"].font = copy.copy(ins["A13"].font)

rub = wb["Rubric"]
for i, row in enumerate(C.RUBRIC[1:], start=2):
    assert rub.cell(i, 1).value == row[0], (rub.cell(i, 1).value, row[0])
    rub.cell(i, 3).value = row[2]

name = "Baseline v4"
if name in wb.sheetnames:
    del wb[name]
ws = wb.create_sheet(name, index=wb.sheetnames.index("Rubric") + 1)
src = wb["Team 2"]
yellow = PatternFill("solid", fgColor="FFF2CC")
ws["A1"] = "NEW (v4) baseline items: enter Y or N per team"
ws["A1"].font = copy.copy(src["A1"].font)
header = ["Baseline item", "Scored under"] + [f"Team {n}" for n in range(1, 9)]
items = [("Modernisation step (characterisation tests first, reviewed PRs)", "Engineering quality"),
         ("AGENTS.md in the team repo", "Engineering quality"),
         ("Security checklist: OWASP Top 10 for Agentic Applications", "Security, guardrails and responsible AI"),
         ("Vulnerability report (scans and red-team, fixed and verified)", "Security, guardrails and responsible AI"),
         ("Token and cost report (placeholder prices labelled)", "Deployment and operations")]
for c, h in enumerate(header, 1):
    cell = ws.cell(3, c, h)
    cell.font = copy.copy(src["A5"].font)
    cell.fill = copy.copy(src["A5"].fill)
for r, (item, crit) in enumerate(items, 4):
    ws.cell(r, 1, item)
    ws.cell(r, 2, crit)
    for c in range(3, 11):
        ws.cell(r, c).fill = yellow
ws.cell(10, 1, "Items present")
for c in range(3, 11):
    col = openpyxl.utils.get_column_letter(c)
    ws.cell(10, c, f'=COUNTIF({col}4:{col}8,"Y")&"/5"')
ws.column_dimensions["A"].width = 62
ws.column_dimensions["B"].width = 38
for c in range(3, 11):
    ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 9
wb.save(PATH)
print("Scoring sheet updated.")
